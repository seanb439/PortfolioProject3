from docx import Document

# Extract content from project specification
doc1 = Document("AI_Portfolio_Final_Project.docx")
print("=== AI_Portfolio_Final_Project.docx ===\n")
for para in doc1.paragraphs:
    if para.text.strip():
        print(para.text)
for table in doc1.tables:
    for row in table.rows:
        for cell in row.cells:
            if cell.text.strip():
                print(cell.text)

print("\n\n=== Sean Burke Resume.docx ===\n")
# Extract content from resume
doc2 = Document("Sean Burke Resume.docx")
for para in doc2.paragraphs:
    if para.text.strip():
        print(para.text)
for table in doc2.tables:
    for row in table.rows:
        for cell in row.cells:
            if cell.text.strip():
                print(cell.text)
