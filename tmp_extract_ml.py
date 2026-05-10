from docx import Document
path = r"c:\Users\seanb\Documents\AIPrinciples\PortfolioProject3\Machine Learning.docx"
doc = Document(path)
for para in doc.paragraphs:
    text = para.text.strip()
    if text:
        print(text)
for table in doc.tables:
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
        if cells:
            print(' | '.join(cells))
